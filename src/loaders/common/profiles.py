"""Validate a DDS instance against a profile, and strip one profile's extensions.

A combined Define-XML + CRF DDS instance carries the extension slots of *both* profiles.
Each profile's JSON Schema is closed, so neither validates the combined instance directly.
The rule (CRF_GEN_PLAN.md §3.4) is **strip-and-validate**: validate against each profile
after removing the *other* profile's extension slots. There is no union profile.
"""
from __future__ import annotations

import copy
import json
import logging
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)

EXTENSION_ANNOTATION = "dds.profile.extension"


def profile_extension_slots(profile_yaml: str | Path) -> set[str]:
    """Return the slot names a profile declares as extensions.

    Reads the differential ``profile.yaml`` and collects every entry under a class's
    ``attributes:`` annotated ``dds.profile.extension: true``.
    """
    import yaml

    with open(profile_yaml, "r", encoding="utf-8") as f:
        profile = yaml.safe_load(f) or {}

    slots: set[str] = set()
    for class_def in (profile.get("classes") or {}).values():
        for slot_name, slot_def in (class_def.get("attributes") or {}).items():
            annotations = (slot_def or {}).get("annotations") or {}
            if annotations.get(EXTENSION_ANNOTATION) is True:
                slots.add(slot_name)
    return slots


def strip_extensions(instance: dict[str, Any], slots: set[str]) -> dict[str, Any]:
    """Return a deep copy of ``instance`` with the named slots removed everywhere.

    :param instance: a DDS MetaDataVersion dict
    :param slots: slot names to remove (from :func:`profile_extension_slots`)
    """
    return _strip(copy.deepcopy(instance), slots)


def _strip(node: Any, slots: set[str]) -> Any:
    if isinstance(node, dict):
        return {k: _strip(v, slots) for k, v in node.items() if k not in slots}
    if isinstance(node, list):
        return [_strip(v, slots) for v in node]
    return node


def validate_against_profile(instance: dict[str, Any], schema_file: str | Path) -> list[str]:
    """Validate a DDS instance against a profile's JSON Schema.

    :param instance: a DDS MetaDataVersion dict
    :param schema_file: path to the profile's ``*.schema.json``
    :return: list of formatted error strings; empty means valid
    """
    import jsonschema

    with open(schema_file, "r", encoding="utf-8") as f:
        schema = json.load(f)

    validator = jsonschema.Draft202012Validator(schema)
    errors = []
    for err in sorted(validator.iter_errors(instance), key=lambda e: list(e.absolute_path)):
        location = "/".join(str(p) for p in err.absolute_path) or "(root)"
        errors.append(f"{location}: {err.message}")
    return errors


def write_validation_excel(errors: list[str], excel_path: str | Path,
                           instance: dict[str, Any], instance_path: str,
                           schema_path: str) -> bool:
    """Write the two-sheet xlsx validation report.

    Mirrors the shape of ``create_define_json._write_validation_excel`` (a Summary sheet
    and a Validation Errors sheet) so both loaders produce interchangeable reports.
    pandas and openpyxl are imported lazily and stay optional dependencies.

    :return: True when the report was written
    """
    try:
        import pandas as pd
        from datetime import datetime
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        logger.warning("pandas/openpyxl not installed - skipping xlsx validation report")
        print("NOTE: install pandas and openpyxl to write the xlsx validation report")
        return False

    passed = not errors
    item_groups = instance.get("itemGroups") or []
    summary = {
        "Validation Run": [datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
        "Schema File": [str(schema_path)],
        "Data File": [str(instance_path)],
        "Status": ["PASSED" if passed else "FAILED"],
        "Total Errors": [len(errors)],
        "Total ItemGroups": [len(item_groups)],
        "Total CodeLists": [len(instance.get("codeLists") or [])],
        "Total StudyEvents": [len(instance.get("studyEvents") or [])],
        "Form ItemGroups": [sum(1 for g in item_groups if g.get("type") == "Form")],
    }
    rows = [{"Error #": i, "Location": e.split(": ", 1)[0], "Issue": e.split(": ", 1)[-1]}
            for i, e in enumerate(errors, start=1)]
    if not rows:
        rows = [{"Error #": "", "Location": "", "Issue": "No validation errors"}]

    try:
        Path(excel_path).parent.mkdir(parents=True, exist_ok=True)
        with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
            pd.DataFrame(summary).to_excel(writer, sheet_name="Summary", index=False)
            pd.DataFrame(rows).to_excel(writer, sheet_name="Validation Errors", index=False)
            for sheet, color in (("Summary", "4472C4"), ("Validation Errors", "C65911")):
                ws = writer.sheets[sheet]
                for cell in ws[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
    except Exception as exc:  # noqa: BLE001 - reporting must never fail the run
        logger.error("could not write validation report %s: %s", excel_path, exc)
        return False
    return True
